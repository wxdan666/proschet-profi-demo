"""
Формирование отчётов и экспорт в Excel.

На вход подаётся pandas.DataFrame (или несколько), на выходе - байтовая
строка с XLSX-файлом для кнопки скачивания. Над таблицей печатается
шапка с названием организации и отчёта, как в бумажных формах.
"""
from datetime import date
from io import BytesIO
from typing import Dict, List, Optional

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from config import CONFIG
from logger import log

_ORG = CONFIG["app"].get("organization", "")
_THIN = Side(style="thin", color="9AA0A6")
_HEAD_FILL = PatternFill("solid", fgColor="E8620C")


def _format_sheet(ws, df: pd.DataFrame, title: str, start_row: int) -> None:
    ws.cell(row=1, column=1, value=_ORG).font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=title).font = Font(bold=True, size=13)
    ws.cell(row=3, column=1,
            value=f"Сформировано {date.today():%d.%m.%Y}").font = Font(
        italic=True, size=9, color="666666")
    for i, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        width = max([len(str(v)) for v in df[col].astype(str)] + [len(str(col))])
        ws.column_dimensions[cell.column_letter].width = min(width + 3, 60)
        for r in range(start_row, start_row + len(df) + 1):
            ws.cell(row=r, column=i).border = Border(
                left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def sheets_to_xlsx(sheets: Dict[str, pd.DataFrame],
                   titles: Optional[Dict[str, str]] = None) -> bytes:
    """Записать несколько таблиц на отдельные листы одной книги."""
    titles = titles or {}
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            sheet = name[:31]
            df.to_excel(writer, sheet_name=sheet, index=False, startrow=4)
            _format_sheet(writer.sheets[sheet], df, titles.get(name, name), 5)
    log.info("Сформирован XLSX: %s", ", ".join(sheets))
    return buf.getvalue()


def dataframe_to_xlsx(df: pd.DataFrame, sheet_name: str = "Отчёт",
                      title: Optional[str] = None) -> bytes:
    """Одна таблица - один лист."""
    return sheets_to_xlsx({sheet_name: df}, {sheet_name: title or sheet_name})


def records_to_xlsx(records: List[Dict], sheet_name: str = "Отчёт",
                    title: Optional[str] = None) -> bytes:
    return dataframe_to_xlsx(pd.DataFrame(records), sheet_name, title)


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
